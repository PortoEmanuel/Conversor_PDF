from django.shortcuts import render, redirect, get_object_or_404
from .forms import ArquivoForm
from .models import Arquivo
from io import BytesIO
from reportlab.pdfgen import canvas
from django.http import FileResponse
from django.contrib import messages
import docx
import openpyxl
from odf.opendocument import load
from odf.text import H, P
import logging
import uuid

logger = logging.getLogger(__name__)
logger.setLevel(logging.ERROR)

def converter(request):
    if request.method == 'POST':
        form = ArquivoForm(request.POST, request.FILES)
        if form.is_valid():
            arquivo = form.save()

            try:
                buffer = BytesIO()
                pdf = canvas.Canvas(buffer)

                if arquivo.arquivo.name.endswith('.docx'):
                    doc = docx.Document(arquivo.arquivo)
                    y_position = 750
                    for para in doc.paragraphs:
                        pdf.drawString(100, y_position, para.text)
                        y_position -= 15
                        if y_position < 50:
                            pdf.showPage()
                            y_position = 750
                elif arquivo.arquivo.name.endswith('.odt'):
                    doc = load(arquivo.arquivo)
                    y_position = 750
                    for elem in doc.getElementsByType(P):
                        pdf.drawString(100, y_position, elem.firstChild.data)
                        y_position -= 15
                        if y_position < 50:
                            pdf.showPage()
                            y_position = 750
                elif arquivo.arquivo.name.endswith('.xlsx'):
                    wb = openpyxl.load_workbook(arquivo.arquivo)
                    sheet = wb.active
                    y_position = 750
                    for row in sheet.iter_rows(values_only=True):
                        row_text = ' '.join([str(cell) if cell is not None else '' for cell in row])
                        pdf.drawString(100, y_position, row_text)
                        y_position -= 15
                        if y_position < 50:
                            pdf.showPage()
                            y_position = 750

                pdf.save()
                buffer.seek(0)
                unique_name = f"{uuid.uuid4().hex}.pdf"
                arquivo.pdf.save(unique_name, buffer, save=True)
            except Exception as e:
                messages.error(request, f"Erro ao converter o arquivo: {e}")
                logger.error(f"Erro ao converter o arquivo: {e}")
                return redirect('converter')

            return redirect('detalhes', pk=arquivo.pk)
    else:
        form = ArquivoForm()

    return render(request, 'conversor/converter.html', {'form': form})

def detalhes(request, pk):
    arquivo = get_object_or_404(Arquivo, pk=pk)

    if request.method == 'POST' and 'download' in request.POST:
        response = FileResponse(open(arquivo.pdf.path, 'rb'), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{arquivo.pdf.name}"'
        return response

    context = {'arquivo': arquivo}
    return render(request, 'conversor/detalhes.html', context)
